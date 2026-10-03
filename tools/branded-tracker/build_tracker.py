import argparse
import datetime as dt
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image as PILImage

ROWS = 500
MONEY = '$#,##0.00;[Red]($#,##0.00);"-"'
DATE = "mm/dd/yyyy"
BUDGET_TYPES = ["Operational", "Event", "Conference", "Collaboration", "Start-Up"]
SOURCES = ["WSA", "CEAS", "Other"]
SEMESTERS = ["Fall", "Spring", "Summer"]
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
STATUSES = ["Pending", "Approved", "Purchased", "Reimbursed", "Declined"]
ROUTES = ["Reimbursement", "OSE Card", "CEAS Card", "Vendor Payment"]
YESNO = ["Yes", "No"]
CATEGORIES = ["Electronics", "Equipment", "Supplies", "Clothing / Merch", "Promo Items", "Other"]
CONDITIONS = ["New", "Good", "Fair", "Needs Repair", "Broken", "Lost"]


def tint(hex_color, amount):
    r, g, b = (int(hex_color[i:i + 2], 16) for i in (0, 2, 4))
    mix = lambda c: round(c + (255 - c) * amount)
    return f"{mix(r):02X}{mix(g):02X}{mix(b):02X}"


class Brand:
    def __init__(self, org_dir):
        cfg = yaml.safe_load((org_dir / "brand.yaml").read_text())
        self.name = cfg["name"]
        self.short = cfg.get("short_name") or cfg["name"]
        self.logo = org_dir / cfg["logo"] if cfg.get("logo") else None
        c = cfg["colors"]
        self.primary = c["primary"].upper()
        self.accent = c["accent"].upper()
        self.on_primary = c.get("on_primary", "FFFFFF").upper()
        self.on_accent = c.get("on_accent", "FFFFFF").upper()
        self.light = tint(self.primary, 0.90)
        self.accent_light = tint(self.accent, 0.75)
        self.font = cfg.get("font") or "Arial"
        self.members = cfg.get("members") or []

    def f(self, size=10, bold=False, color="222222", italic=False):
        return Font(name=self.font, size=size, bold=bold, color=color, italic=italic)


def build(org_dir, out_path, fiscal_year, fall_end, spring_end):
    b = Brand(org_dir)
    thin = Side(style="thin", color=tint(b.primary, 0.7))
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor=b.primary)
    calc_fill = PatternFill("solid", fgColor="F2F2F2")
    light_fill = PatternFill("solid", fgColor=b.light)
    accent_fill = PatternFill("solid", fgColor=b.accent_light)
    ex_font = b.f(italic=True, color="8C8C8C")

    wb = Workbook()

    def header(ws, row, cols):
        for i, (name, width) in enumerate(cols, 1):
            c = ws.cell(row=row, column=i, value=name)
            c.font = b.f(10, True, b.on_primary)
            c.fill = head_fill
            c.border = box
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[L(i)].width = width
        ws.row_dimensions[row].height = 30

    def grid(ws, first, last, ncols, calc_cols=(), fmt=None):
        fmt = fmt or {}
        for r in range(first, last + 1):
            for c in range(1, ncols + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = box
                cell.font = b.f()
                if c in calc_cols:
                    cell.fill = calc_fill
                if c in fmt:
                    cell.number_format = fmt[c]

    def title(ws, text, sub=None):
        ws["A1"] = text
        ws["A1"].font = b.f(16, True, b.primary)
        if sub:
            ws["A2"] = sub
            ws["A2"].font = b.f(10, italic=True, color="595959")

    def dv(ws, formula, rng, strict=True):
        d = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
        ws.add_data_validation(d)
        d.add(rng)

    lists = wb.active
    lists.title = "Lists"
    for ci, (head, vals) in enumerate([("Budget Type", BUDGET_TYPES), ("Source", SOURCES), ("Semester", SEMESTERS),
                                       ("Month", MONTHS), ("Status", STATUSES), ("Payment Route", ROUTES), ("Yes/No", YESNO), ("Category", CATEGORIES), ("Condition", CONDITIONS)], 1):
        lists.cell(row=1, column=ci, value=head).font = b.f(10, True)
        lists.column_dimensions[L(ci)].width = 18
        for ri, v in enumerate(vals, 2):
            lists.cell(row=ri, column=ci, value=v).font = b.f()
    rng = lambda col, n: f"=Lists!${col}$2:${col}${n + 1}"

    st = wb.create_sheet("Settings")
    title(st, "Settings", "Yellow cells are yours to edit. Everything else in the workbook reads from here.")
    st.column_dimensions["A"].width = 34
    st.column_dimensions["B"].width = 18
    st.column_dimensions["C"].width = 70
    rows = [
        (4, "Club name", b.name, None, "From the club's brand file."),
        (5, "Fiscal year", fiscal_year, None, "WSA fiscal year runs July 1 to June 30 (Bylaws F26 §5.02(a))."),
        (7, "WSA yearly cap", 6500, MONEY, "Per RSO per fiscal year, all categories except collaboration (Bylaws F26 §6.03(a))."),
        (8, "WSA collaboration cap", 1500, MONEY, "Separate from the $6,500 cap (Step 2, Art. 6.03; rule COL-04)."),
        (9, "WSA start-up cap", 500, MONEY, "Only for RSOs established one semester or less (Step 2, Art. 6.03; rule OTH-01)."),
        (10, "CEAS budget", 0, MONEY, "Enter the amount CEAS approved, if the club has CEAS funding."),
        (12, "Fall semester ends", fall_end, DATE, "Operational money from fall expires this day. WMU 2026-27 calendar: Dec. 19."),
        (13, "Spring semester ends", spring_end, DATE, "Operational money from spring expires this day. WMU 2026-27 calendar: May 1."),
        (14, "Academic year ends", "=B13", DATE, "Start-up money expires at the end of the academic year (rule OTH-01)."),
        (15, "Inventory check every (days)", 120, "0", "How often each inventory item should be checked. 120 days is about once a semester."),
    ]
    yellow = PatternFill("solid", fgColor="FFF4C2")
    for r, label, val, fmt, note in rows:
        st.cell(row=r, column=1, value=label).font = b.f(10, True)
        c = st.cell(row=r, column=2, value=val)
        c.font = b.f()
        c.border = box
        if fmt:
            c.number_format = fmt
        if not (isinstance(val, str) and val.startswith("=")):
            c.fill = yellow
        st.cell(row=r, column=3, value=note).font = b.f(9, italic=True, color="595959")
    st["A6"], st["A11"], st["A16"] = "Caps", "Expiration dates", "Members (used in the Purchased By dropdown)"
    for a in ("A6", "A11", "A16"):
        st[a].font = b.f(11, True, b.primary)
    for i in range(30):
        c = st.cell(row=17 + i, column=1, value=b.members[i] if i < len(b.members) else None)
        c.fill, c.border, c.font = yellow, box, b.f()

    fd = wb.create_sheet("Funding")
    title(fd, "Funding", "One row per item on a funding letter. Grey columns calculate on their own.")
    fcols = [("FID", 9), ("Item (as on the funding letter)", 30), ("Budget Type", 15), ("Source", 10), ("Semester", 10),
             ("Deliberation Month", 13), ("Event / Conference Date", 13), ("Amount Approved", 14), ("Expires On", 13),
             ("Spent", 13), ("Remaining", 13), ("Status", 15), ("Funding Letter (file name)", 26), ("Notes", 34), ("WSA #", 7)]
    HR = 4
    header(fd, HR, fcols)
    first, last = HR + 1, HR + ROWS
    for r in range(first, last + 1):
        fd[f"I{r}"] = (f'=IF(A{r}="","",IF(OR(C{r}="Event",C{r}="Collaboration",C{r}="Conference"),IF(G{r}="","Add date",G{r}),'
                       f'IF(C{r}="Operational",IF(E{r}="Fall",Settings!$B$12,IF(E{r}="Spring",Settings!$B$13,"Add semester")),'
                       f'IF(C{r}="Start-Up",Settings!$B$14,""))))')
        fd[f"J{r}"] = f'=IF(A{r}="","",SUMIFS(Purchases!$I:$I,Purchases!$B:$B,A{r},Purchases!$F:$F,"<>Declined"))'
        fd[f"K{r}"] = f'=IF(A{r}="","",H{r}-J{r})'
        fd[f"L{r}"] = (f'=IF(A{r}="","",IF(K{r}<0,"Over budget",IF(K{r}=0,"Fully used",IF(ISNUMBER(I{r}),'
                       f'IF(I{r}<TODAY(),"Expired",IF(I{r}-TODAY()<=14,"Expiring soon","Open")),"Open"))))')
        fd[f"O{r}"] = f'=IF(AND(D{r}="WSA",LEFT(A{r},3)<>"EX-"),COUNTIFS(D${first}:D{r},"WSA",A${first}:A{r},"<>EX-*"),"")'
    grid(fd, first, last, len(fcols), calc_cols=(9, 10, 11, 12, 15), fmt={7: DATE, 8: MONEY, 9: DATE, 10: MONEY, 11: MONEY})
    fd.column_dimensions["O"].hidden = True
    dv(fd, rng("A", 5), f"C{first}:C{last}")
    dv(fd, rng("B", 3), f"D{first}:D{last}")
    dv(fd, rng("C", 3), f"E{first}:E{last}")
    dv(fd, rng("D", 12), f"F{first}:F{last}")
    for txt, col in [("Over budget", "C00000"), ("Expired", "7F7F7F"), ("Expiring soon", "C55A11")]:
        fd.conditional_formatting.add(f"L{first}:L{last}", FormulaRule(formula=[f'$L{first}="{txt}"'], font=Font(name=b.font, bold=True, color=col)))
    fd.freeze_panes = f"B{first}"

    pu = wb.create_sheet("Purchases")
    title(pu, "Purchases", "One row per receipt. Pick the FID from the Funding tab and the item fills in.")
    pcols = [("PID", 8), ("FID", 9), ("Item", 28), ("Budget Type", 14), ("Source", 10), ("Status", 13), ("Purchased By", 16),
             ("Date", 12), ("Amount", 13), ("Payment Route", 16), ("Distribution Form Sent?", 13), ("Receipt File", 26),
             ("Receipt #", 9), ("Notes", 30), ("Track in Inventory?", 12), ("In Inventory", 14)]
    header(pu, HR, pcols)
    for r in range(first, last + 1):
        for col, src in (("C", "B"), ("D", "C"), ("E", "D")):
            pu[f"{col}{r}"] = f'=IF(B{r}="","",IFERROR(INDEX(Funding!${src}:${src},MATCH(B{r},Funding!$A:$A,0)),"Check FID"))'
        pu[f"M{r}"] = f'=IF(OR(B{r}="",F{r}="Declined"),"",COUNTIFS(B${first}:B{r},B{r},F${first}:F{r},"<>Declined"))'
        pu[f"P{r}"] = f'=IF(OR(A{r}="",O{r}<>"Yes"),"",IF(COUNTIF(Inventory!$D:$D,A{r})>0,"Listed","Not listed yet"))'
    grid(pu, first, last, len(pcols), calc_cols=(3, 4, 5, 13, 16), fmt={8: DATE, 9: MONEY})
    dv(pu, rng("G", 2), f"O{first}:O{last}")
    pu.conditional_formatting.add(f"P{first}:P{last}", FormulaRule(formula=[f'$P{first}="Not listed yet"'], font=Font(name=b.font, bold=True, color="C55A11")))
    dv(pu, f"=Funding!$A${first}:$A${last}", f"B{first}:B{last}")
    dv(pu, rng("E", 5), f"F{first}:F{last}")
    dv(pu, "=Settings!$A$17:$A$46", f"G{first}:G{last}", strict=False)
    dv(pu, rng("F", 4), f"J{first}:J{last}")
    dv(pu, rng("G", 2), f"K{first}:K{last}")
    pu.conditional_formatting.add(f"K{first}:K{last}", FormulaRule(formula=[f'AND($K{first}="No",OR($F{first}="Purchased",$F{first}="Reimbursed"))'], fill=PatternFill("solid", fgColor="FBE2D5")))
    pu.freeze_panes = f"B{first}"

    ex_f = [
        ["EX-01", "Example: Monitors", "Operational", "WSA", "Fall", "September", None, 300, None, None, None, None, "Example letter.pdf", "Example row, not counted on the dashboard."],
        ["EX-02", "Example: Pizza", "Event", "WSA", "Fall", "October", dt.date(2026, 10, 30), 120, None, None, None, None, "Example letter.pdf", "Food is only allowed on event and collaboration budgets."],
    ]
    for i, row in enumerate(ex_f):
        for c, v in enumerate(row, 1):
            if v is not None:
                fd.cell(row=first + i, column=c, value=v)
        for c in range(1, len(fcols) + 1):
            fd.cell(row=first + i, column=c).font = ex_font
    ex_p = [
        ["EX-P1", "EX-01", None, None, None, "Reimbursed", "Example Member", dt.date(2026, 9, 20), 189.99, "Reimbursement", "Yes", "example-receipt-1.pdf", None, "Example row.", "Yes"],
        ["EX-P2", "EX-01", None, None, None, "Purchased", "Example Member", dt.date(2026, 9, 27), 94.50, "OSE Card", "No", "example-receipt-2.pdf", None, "Distribution form still needed.", "No"],
    ]
    for i, row in enumerate(ex_p):
        for c, v in enumerate(row, 1):
            if v is not None:
                pu.cell(row=first + i, column=c, value=v)
        for c in range(1, len(pcols) + 1):
            pu.cell(row=first + i, column=c).font = ex_font

    wt = wb.create_sheet("WSAAC Tracker")
    title(wt, "WSAAC RSO Budget Tracker", "Same layout as WSAAC's official tracker. Fills in from the Funding and Purchases tabs. WSA items only.")
    wcols = [("Budget Items", 30), ("Amount Approved", 15), ("Amount Remaining", 16), ("Receipt 1", 12), ("Receipt 2", 12),
             ("Receipt 3", 12), ("Receipt 4", 12), ("Receipt 5+", 12), ("FID", 9)]
    header(wt, HR, wcols)
    WN = 100
    for k in range(1, WN + 1):
        r = HR + k
        m = f"MATCH({k},Funding!$O:$O,0)"
        wt[f"I{r}"] = f'=IFERROR(INDEX(Funding!$A:$A,{m}),"")'
        wt[f"A{r}"] = f'=IF(I{r}="","",INDEX(Funding!$B:$B,{m}))'
        wt[f"B{r}"] = f'=IF(I{r}="","",INDEX(Funding!$H:$H,{m}))'
        for j, col in enumerate("DEFG", 1):
            wt[f"{col}{r}"] = f'=IF(I{r}="","",SUMIFS(Purchases!$I:$I,Purchases!$B:$B,I{r},Purchases!$M:$M,{j}))'
        wt[f"H{r}"] = f'=IF(I{r}="","",SUMIFS(Purchases!$I:$I,Purchases!$B:$B,I{r},Purchases!$M:$M,">=5"))'
        wt[f"C{r}"] = f'=IF(I{r}="","",B{r}-SUM(D{r}:H{r}))'
    grid(wt, HR + 1, HR + WN, len(wcols), calc_cols=tuple(range(1, 10)), fmt={2: MONEY, 3: MONEY, 4: MONEY, 5: MONEY, 6: MONEY, 7: MONEY, 8: MONEY})
    tr = HR + WN + 1
    wt[f"A{tr}"], wt[f"C{tr}"] = "Total Remaining", f"=SUM(C{HR + 1}:C{HR + WN})"
    for col in "ABC":
        wt[f"{col}{tr}"].font = b.f(10, True)
        wt[f"{col}{tr}"].fill = accent_fill
        wt[f"{col}{tr}"].border = box
    wt[f"C{tr}"].number_format = MONEY
    wt.column_dimensions["I"].hidden = True
    wt.freeze_panes = f"A{HR + 1}"

    iv = wb.create_sheet("Inventory")
    title(iv, "Inventory Checklist", "Everything the club owns, who has it, and when it was last checked. Link a purchase by its PID and the details fill in.")
    icols = [("Asset ID", 10), ("Item", 28), ("Qty", 6), ("PID", 9), ("FID", 9), ("Source", 10), ("Date Acquired", 12),
             ("Purchase Amount", 13), ("Category", 15), ("Stored At", 20), ("Checked Out To", 16), ("Checked Out On", 13),
             ("Condition", 13), ("Last Checked", 12), ("Check Status", 14), ("Notes", 30)]
    header(iv, HR, icols)
    for r in range(first, last + 1):
        iv[f"E{r}"] = f'=IF(D{r}="","",IFERROR(INDEX(Purchases!$B:$B,MATCH(D{r},Purchases!$A:$A,0)),"Check PID"))'
        iv[f"F{r}"] = f'=IF(OR(E{r}="",E{r}="Check PID"),"",IFERROR(INDEX(Funding!$D:$D,MATCH(E{r},Funding!$A:$A,0)),""))'
        iv[f"G{r}"] = f'=IF(D{r}="","",IFERROR(INDEX(Purchases!$H:$H,MATCH(D{r},Purchases!$A:$A,0)),""))'
        iv[f"H{r}"] = f'=IF(D{r}="","",IFERROR(INDEX(Purchases!$I:$I,MATCH(D{r},Purchases!$A:$A,0)),""))'
        iv[f"O{r}"] = (f'=IF(A{r}="","",IF(OR(M{r}="Lost",M{r}="Broken",M{r}="Needs Repair"),"Needs action",'
                       f'IF(N{r}="","Never checked",IF(TODAY()-N{r}>Settings!$B$15,"Check due","OK"))))')
    grid(iv, first, last, len(icols), calc_cols=(5, 6, 7, 8, 15), fmt={7: DATE, 8: MONEY, 12: DATE, 14: DATE})
    dv(iv, f"=Purchases!$A${first}:$A${last}", f"D{first}:D{last}", strict=False)
    dv(iv, rng("H", len(CATEGORIES)), f"I{first}:I{last}")
    dv(iv, "=Settings!$A$17:$A$46", f"K{first}:K{last}", strict=False)
    dv(iv, rng("I", len(CONDITIONS)), f"M{first}:M{last}")
    for txt, col in [("Needs action", "C00000"), ("Check due", "C55A11"), ("Never checked", "C55A11"), ("OK", "548235")]:
        iv.conditional_formatting.add(f"O{first}:O{last}", FormulaRule(formula=[f'$O{first}="{txt}"'], font=Font(name=b.font, bold=True, color=col)))
    ex_i = ["EX-INV1", "Example: Monitors", 2, "EX-P1", None, None, None, None, "Electronics", "Example: club storage room", None, None, "Good", dt.date(2026, 9, 25), None, "Example row, not counted on the dashboard."]
    for c, v in enumerate(ex_i, 1):
        if v is not None:
            iv.cell(row=first, column=c, value=v)
    for c in range(1, len(icols) + 1):
        iv.cell(row=first, column=c).font = ex_font
    iv.freeze_panes = f"C{first}"

    db = wb.create_sheet("Dashboard", 0)
    db.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", (3, 36, 14, 14, 14, 12, 11, 26)):
        db.column_dimensions[col].width = w
    for r in range(1, 6):
        db.row_dimensions[r].height = 20
    db["B2"] = f"{b.name}"
    db["B2"].font = b.f(18, True, b.primary)
    db["B3"] = '="Finance Tracker  |  Fiscal Year "&Settings!B5'
    db["B3"].font = b.f(11, color="595959")
    db["B4"] = '="Updated "&TEXT(TODAY(),"mmmm d, yyyy")'
    db["B4"].font = b.f(9, italic=True, color="8C8C8C")
    for c in "BCDEFGH":
        db[f"{c}6"].fill = PatternFill("solid", fgColor=b.accent)
    db.row_dimensions[6].height = 5

    if b.logo and b.logo.exists():
        with PILImage.open(b.logo) as im:
            w, h = im.size
        img = XLImage(str(b.logo))
        target_h = 95
        img.height = target_h
        img.width = round(w * target_h / h)
        if img.width > 260:
            img.width = 260
            img.height = round(h * 260 / w)
        db.add_image(img, "G1")

    db["B8"] = "Funding by source"
    db["B8"].font = b.f(12, True, b.primary)
    dh = ["Source", "Approved", "Spent", "Remaining", "Cap", "Cap Used", "Cap Used (bar)"]
    for i, t in enumerate(dh, 2):
        c = db.cell(row=9, column=i, value=t)
        c.font, c.fill, c.border = b.f(10, True, b.on_primary), head_fill, box
    srcs = [
        ("WSA (Operational, Event, Conference)", '"WSA"', ['Funding!$C:$C,"<>Collaboration"', 'Funding!$C:$C,"<>Start-Up"'], "Settings!$B$7"),
        ("WSA Collaboration", '"WSA"', ['Funding!$C:$C,"Collaboration"'], "Settings!$B$8"),
        ("WSA Start-Up", '"WSA"', ['Funding!$C:$C,"Start-Up"'], "Settings!$B$9"),
        ("CEAS", '"CEAS"', [], "Settings!$B$10"),
        ("Other", '"Other"', [], None),
    ]
    for i, (label, src, extra, cap) in enumerate(srcs):
        r = 10 + i
        crit = ",".join([f"Funding!$D:$D,{src}", 'Funding!$A:$A,"<>EX-*"'] + extra)
        db[f"B{r}"] = label
        db[f"C{r}"] = f"=SUMIFS(Funding!$H:$H,{crit})"
        db[f"D{r}"] = f"=SUMIFS(Funding!$J:$J,{crit})"
        db[f"E{r}"] = f"=C{r}-D{r}"
        db[f"F{r}"] = f"={cap}" if cap else "-"
        db[f"G{r}"] = f'=IF(OR(F{r}="-",N(F{r})=0),"-",C{r}/F{r})'
        db[f"H{r}"] = f'=IF(G{r}="-","",REPT("█",MIN(20,ROUND(G{r}*20,0)))&REPT("░",20-MIN(20,ROUND(G{r}*20,0))))'
        for col in "BCDEFGH":
            cell = db[f"{col}{r}"]
            cell.border = box
            cell.font = b.f(10, col == "B")
            if i % 2:
                cell.fill = light_fill
        for col in "CDEF":
            db[f"{col}{r}"].number_format = MONEY
        db[f"G{r}"].number_format = "0%"
        db[f"H{r}"].font = Font(name=b.font, size=10, color=b.primary)
        db.conditional_formatting.add(f"G{r}", FormulaRule(formula=[f'AND(ISNUMBER(G{r}),G{r}>1)'], font=Font(name=b.font, bold=True, color="C00000")))
    tr = 10 + len(srcs)
    db[f"B{tr}"] = "Total"
    for col in "CDE":
        db[f"{col}{tr}"] = f"=SUM({col}10:{col}{tr - 1})"
        db[f"{col}{tr}"].number_format = MONEY
    for col in "BCDEFGH":
        cell = db[f"{col}{tr}"]
        cell.font, cell.fill, cell.border = b.f(10, True, b.on_accent), PatternFill("solid", fgColor=b.accent), box

    ar = tr + 3
    db[f"B{ar}"] = "Needs attention"
    db[f"B{ar}"].font = b.f(12, True, b.primary)
    F_ = f"Funding!$A${first}:$A${last}"
    alerts = [
        ("Items expiring in the next 14 days", f'=COUNTIFS(Funding!$L:$L,"Expiring soon",Funding!$A:$A,"<>EX-*")', "Spend it or lose it. Expired money cannot be reimbursed."),
        ("Items expired with money left", f'=COUNTIFS(Funding!$L:$L,"Expired",Funding!$A:$A,"<>EX-*")', "Rule MON-06 / PAY-11."),
        ("Items over budget", f'=COUNTIFS(Funding!$L:$L,"Over budget",Funding!$A:$A,"<>EX-*")', "Shifting money between items needs WSAAC approval first (OPS-07)."),
        ("Purchases waiting on reimbursement", f'=COUNTIFS(Purchases!$F:$F,"Purchased",Purchases!$J:$J,"Reimbursement",Purchases!$A:$A,"<>EX-*")', "Submit receipts to OSE promptly."),
        ("Distribution forms not sent", f'=COUNTIFS(Purchases!$K:$K,"No",Purchases!$F:$F,"<>Declined",Purchases!$A:$A,"<>EX-*")', "Every payment starts with the Allocation Distribution Request Form."),
        ("Category caps exceeded", f'=COUNTIFS(G10:G{tr - 1},">1")', "Approved total is above the cap in Settings."),
        ("Purchases to add to inventory", '=COUNTIFS(Purchases!$P:$P,"Not listed yet",Purchases!$A:$A,"<>EX-*")', "Marked Track in Inventory? but not on the Inventory tab yet."),
        ("Inventory items due for a check", '=COUNTIFS(Inventory!$O:$O,"Check due",Inventory!$A:$A,"<>EX-*")+COUNTIFS(Inventory!$O:$O,"Never checked",Inventory!$A:$A,"<>EX-*")', "Find it, check its condition, and update Last Checked."),
        ("Inventory lost, broken, or needs repair", '=COUNTIFS(Inventory!$O:$O,"Needs action",Inventory!$A:$A,"<>EX-*")', "Decide whether to repair, replace, or write it off."),
    ]
    for i, (label, formula, note) in enumerate(alerts):
        r = ar + 1 + i
        db[f"B{r}"], db[f"C{r}"], db[f"D{r}"] = label, formula, note
        db[f"B{r}"].font = b.f()
        db[f"C{r}"].font = b.f(11, True, b.primary)
        db[f"C{r}"].alignment = Alignment(horizontal="center")
        db[f"D{r}"].font = b.f(9, italic=True, color="595959")
        db[f"B{r}"].border = db[f"C{r}"].border = box
        db.conditional_formatting.add(f"C{r}", FormulaRule(formula=[f"C{r}>0"], fill=PatternFill("solid", fgColor="FBE2D5"), font=Font(name=b.font, bold=True, color="C00000")))

    ch = BarChart()
    ch.type = "bar"
    ch.style = 10
    ch.title = "Approved vs. spent"
    ch.y_axis.numFmt = "$#,##0"
    ch.y_axis.majorGridlines = None
    data = Reference(db, min_col=3, max_col=4, min_row=9, max_row=tr - 1)
    cats = Reference(db, min_col=2, min_row=10, max_row=tr - 1)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.series[0].graphicalProperties.solidFill = b.primary
    ch.series[0].graphicalProperties.line.solidFill = b.primary
    ch.series[1].graphicalProperties.solidFill = b.accent
    ch.series[1].graphicalProperties.line.solidFill = b.accent
    ch.height, ch.width = 7.5, 17
    db.add_chart(ch, f"B{ar + len(alerts) + 3}")

    hw = wb.create_sheet("How to Use", 1)
    hw.sheet_view.showGridLines = False
    hw.column_dimensions["A"].width = 3
    hw.column_dimensions["B"].width = 105
    lines = [
        (f"How to use the {b.short} Finance Tracker", b.f(16, True, b.primary)),
        ("", None),
        ("Setup (once a year)", b.f(12, True, b.primary)),
        ("1. Settings tab: check the fiscal year, caps, and semester end dates. Add your members' names.", b.f()),
        ("2. Delete the grey EX- example rows on Funding and Purchases when you are ready (they are already left out of the dashboard).", b.f()),
        ("", None),
        ("When a funding letter arrives", b.f(12, True, b.primary)),
        ("3. Funding tab: add one row per approved item. Give it the next FID (F-001, F-002, ...). Type the item name exactly as it appears on the letter, because OSE matches receipts to it.", b.f()),
        ("4. For events, collaborations and conferences, enter the event date. For operational items, pick the semester. The Expires On date fills in by itself.", b.f()),
        ("", None),
        ("Every time something is bought", b.f(12, True, b.primary)),
        ("5. Purchases tab: one row per receipt. Pick the FID and the item, budget type and source fill in. Save the receipt file in the club's drive folder and type its file name.", b.f()),
        ("6. Update the Status as it moves along: Pending, Approved, Purchased, Reimbursed (or Declined).", b.f()),
        ("7. Mark Distribution Form Sent? as Yes once the Allocation Distribution Request Form is submitted to OSE.", b.f()),
        ("", None),
        ("What you get", b.f(12, True, b.primary)),
        ("Dashboard: approved, spent and remaining money by source, how much of each cap is used, and anything that needs attention.", b.f()),
        ("WSAAC Tracker: WSAAC's official RSO Budget Tracker layout, filled in for you. Use it if WSAAC or OSE asks for the tracker.", b.f()),
        ("Inventory: what the club owns, who has it, and what is due for a check. The Dashboard flags anything missing or overdue.", b.f()),
        ("", None),
        ("Inventory checklist", b.f(12, True, b.primary)),
        ("8. When you buy something the club will keep (monitors, cables, a banner), set Track in Inventory? to Yes on the Purchases tab.", b.f()),
        ("9. Inventory tab: add a row, type the PID, and the date, amount and funding source fill in. Then fill in where it is stored and its condition.", b.f()),
        ("10. Lending something out? Fill in Checked Out To and Checked Out On, and clear them when it comes back.", b.f()),
        ("11. Once a semester, walk through the list, check each item, and update Condition and Last Checked. Do this before officers change over.", b.f()),
        ("", None),
        ("Rules worth remembering", b.f(12, True, b.primary)),
        ("No food on Operational or Conference budgets. Food is only for on-campus events, up to $20 per expected attendee.", b.f()),
        ("Nothing bought before the funding letter is covered. Keep every itemized receipt.", b.f()),
        ("Money expires: event and collaboration money after the event, conference money after the trip, operational money at the end of the semester.", b.f()),
        ("To spend different amounts on approved items, the total cannot go over the letter, and WSAAC must approve first (OPS-07).", b.f()),
        ("", None),
        ("Privacy", b.f(12, True, b.primary)),
        ("This file has members' names and receipt details. Keep it in the club's private drive. Never upload it to GitHub or post it publicly.", b.f()),
    ]
    for i, (t, f) in enumerate(lines, 2):
        c = hw.cell(row=i, column=2, value=t)
        if f:
            c.font = f
        c.alignment = Alignment(wrap_text=True, vertical="top")

    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.move_sheet("Lists", offset=10)
    lists.sheet_state = "hidden"
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = b.primary if ws.title in ("Dashboard", "How to Use", "Inventory") else tint(b.primary, 0.5)
    wb.active = 0
    wb.save(out_path)


if __name__ == "__main__":
    p = argparse.ArgumentParser(description="Build a club-branded WSA finance tracker.")
    p.add_argument("org_dir", type=Path)
    p.add_argument("--out", type=Path, required=True)
    p.add_argument("--fiscal-year", default="2026-27")
    p.add_argument("--fall-end", default="2026-12-19")
    p.add_argument("--spring-end", default="2027-05-01")
    a = p.parse_args()
    build(a.org_dir, a.out, a.fiscal_year, dt.date.fromisoformat(a.fall_end), dt.date.fromisoformat(a.spring_end))
    print(f"Built {a.out}")
